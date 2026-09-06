# -*- coding: utf-8 -*-
"""
听评课报告生成脚本（星火个性化版）
基于课堂录音转写文本，生成结构化听评课建议报告
评估标准依据《星火教育 1对1 课堂听评课评价标准》：教学环节(75) + 教学风格(10) + 学生反馈(15)
支持多种输出格式
"""

import os
import sys
import json
from pathlib import Path
from datetime import datetime

# Windows GBK 控制台兼容：stdout/stderr 统一 UTF-8 容错输出，避免非 ASCII 字符崩溃
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

# 可选依赖：Word/PDF 导出
# 安装：pip install python-docx reportlab
try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

try:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib import colors
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    HAS_PDF = True
except ImportError:
    HAS_PDF = False

# 默认报告输出目录（统一位置）
DEFAULT_REPORT_DIR = './reports'


class ClassroomEvaluationReport:
    """听评课报告生成器（星火个性化版）"""

    # 听评课评估维度（依据《星火教育 1对1 课堂听评课评价标准》）
    # weight 即满分分值，总分 100 分
    EVALUATION_DIMENSIONS = {
        '教学环节': {
            'description': '以教师为主导，学生为主体，训练为主线，思维为主攻，能力为主旨',
            'weight': 75,
            'sub_items': [
                {
                    'name': '课堂导入', 'max_score': 5,
                    'indicators': [
                        '导入与内容贴合度高，呼应主题',
                        '导入符合学生的性格兴趣、学习情况'
                    ]
                },
                {
                    'name': '主题细化', 'max_score': 10,
                    'indicators': [
                        '主题选取合理，进行合理的模块分解，题型设置合理',
                        '例题符合学生水平，符合考纲和课程标准，符合教学进度'
                    ]
                },
                {
                    'name': '例题讲解', 'max_score': 40,
                    'indicators': [
                        '根据学生水平进行了合理知识铺垫',
                        '从题干到思路过渡合理',
                        '逻辑缜密，步骤清晰；关注学生',
                        '层层递进、深入浅出，重难点、易错点突出',
                        '强调本题型方法重要性，总结易错点避雷，学生有核心逻辑备注',
                        '目标达成，学生完全掌握解题思路'
                    ]
                },
                {
                    'name': '变式选题', 'max_score': 10,
                    'indicators': [
                        '所选习题新颖、精准且贴合考情；题量满足学生需要',
                        '变式题与变式题之间、与例题之间层次分明，有清晰的逻辑关系'
                    ]
                },
                {
                    'name': '模块小结', 'max_score': 5,
                    'indicators': [
                        '总结课堂内容梳理清晰',
                        '明确技巧方法适用性，突出重难点/易错点',
                        '强调逻辑性'
                    ]
                },
                {
                    'name': '作业布置', 'max_score': 5,
                    'indicators': [
                        '作业符合学生水平，体现本堂课重难点和学生错误',
                        '符合艾宾浩斯遗忘曲线'
                    ]
                },
            ]
        },
        '教学风格': {
            'description': '教师形象与语言表达',
            'weight': 10,
            'sub_items': [
                {
                    'name': '形象', 'max_score': 5,
                    'indicators': [
                        '亲和力强，笑容自然具有感染力',
                        '动作从容，具有强化语言效果',
                        '气场欢快而严肃、可亲又可敬'
                    ]
                },
                {
                    'name': '语言', 'max_score': 5,
                    'indicators': [
                        '语调抑扬顿挫，富有感染力',
                        '语言表达符合学生认知水平和年龄特点，生动形象富有趣味性；表达精练；逻辑用语严谨，具有启发性'
                    ]
                },
            ]
        },
        '学生反馈': {
            'description': '以学生视角为基准，重点关注学生听完是否真正理解',
            'weight': 15,
            'sub_items': [
                {
                    'name': '倾听配合', 'max_score': 5,
                    'indicators': [
                        '学生倾听认真',
                        '课堂过程有记笔记/查阅/回应等辅助行为'
                    ]
                },
                {
                    'name': '师生互动', 'max_score': 5,
                    'indicators': [
                        '师生间有眼神、语言或动作方面对学习有效的互动'
                    ]
                },
                {
                    'name': '课程目标', 'max_score': 5,
                    'indicators': [
                        '学生能充分理解讲课思路并能独立完成习题',
                        '能够达到预设目标'
                    ]
                },
            ]
        }
    }

    def __init__(self, transcript_data, teacher_name=None, course_name=None, grade=None, student_name=None):
        """
        初始化报告生成器

        Args:
            transcript_data: 转写数据（字典或 JSON 文件路径）
            teacher_name: 教师姓名
            course_name: 课程名称
            grade: 年级
            student_name: 学生姓名
        """
        if isinstance(transcript_data, (str, Path)):
            with open(transcript_data, 'r', encoding='utf-8') as f:
                self.transcript = json.load(f)
        else:
            self.transcript = transcript_data

        self.teacher_name = teacher_name or '待填写'
        self.course_name = course_name or '待填写'
        self.grade = grade or '待填写'
        self.student_name = student_name or '待填写'
        self.evaluation_date = datetime.now().strftime('%Y-%m-%d')

    def generate_report(self, format='markdown', output_path=None, custom_template=None):
        """
        生成听评课报告

        Args:
            format: 输出格式 (markdown, json, html, text)
            output_path: 输出文件路径
            custom_template: 自定义评估维度（字典）

        Returns:
            报告内容（字符串）
        """
        dimensions = custom_template or self.EVALUATION_DIMENSIONS
        transcript_text = self.transcript.get('text', '')

        # 分析转写内容
        analysis = self._analyze_transcript(transcript_text, dimensions)

        # 生成报告（Word/PDF 直接在方法内写入文件）
        if format == 'docx':
            if not HAS_DOCX:
                raise RuntimeError('缺少 python-docx，请先安装: pip install python-docx')
            if not output_path:
                raise ValueError('docx/pdf 格式必须指定输出路径')
            return self._generate_docx_report(analysis, dimensions, output_path)
        if format == 'pdf':
            if not HAS_PDF:
                raise RuntimeError('缺少 reportlab，请先安装: pip install reportlab')
            if not output_path:
                raise ValueError('docx/pdf 格式必须指定输出路径')
            return self._generate_pdf_report(analysis, dimensions, output_path)

        # 文本格式生成
        if format == 'markdown':
            report = self._generate_markdown_report(analysis, dimensions)
        elif format == 'json':
            report = self._generate_json_report(analysis, dimensions)
        elif format == 'text':
            report = self._generate_text_report(analysis, dimensions)
        elif format == 'html':
            report = self._generate_html_report(analysis, dimensions)
        else:
            raise ValueError(f"不支持的输出格式: {format}")

        # 保存报告（文本格式）
        if output_path:
            output_path = Path(output_path)
            output_path.parent.mkdir(parents=True, exist_ok=True)
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(report)
            print(f"报告已保存: {output_path}")

        return report

    @staticmethod
    def _iter_dimension_items(dim_info):
        """迭代维度的评价细项：有 sub_items 返回子项，否则返回维度本身"""
        if dim_info.get('sub_items'):
            return dim_info['sub_items']
        return [{'name': dim_info.get('name', ''), 'max_score': dim_info.get('weight', 0),
                 'indicators': dim_info.get('indicators', [])}]

    def _analyze_transcript(self, transcript_text, dimensions):
        """
        分析转写文本（简化版分析逻辑）

        实际应用中，这里应该调用 LLM API 进行深度分析
        当前版本返回示例分析结果（子项按满分 75% 给出示例分，供 AI 覆盖）
        """
        # 简单文本统计
        word_count = len(transcript_text)
        sentence_count = transcript_text.count('。') + transcript_text.count('！') + transcript_text.count('？')

        # 基础分析（实际应由 AI 分析）
        analysis = {
            'basic_info': {
                'word_count': word_count,
                'sentence_count': sentence_count,
                'duration': self.transcript.get('duration', 0)
            },
            'dimensions': {},
            'highlights': [],
            'suggestions': [],
            'overall_score': 0
        }

        # 为每个维度生成示例评估
        # 注意：这里是示例逻辑，实际应该由 AI 分析
        for dim_name, dim_info in dimensions.items():
            dim_score = 0
            sub_scores = {}
            for item in self._iter_dimension_items(dim_info):
                # 示例分数：满分的 75%（取整），供 AI 覆盖
                example_score = round(item['max_score'] * 0.75)
                sub_scores[item['name']] = example_score
                dim_score += example_score

            analysis['dimensions'][dim_name] = {
                'score': dim_score,
                'sub_scores': sub_scores,
                'strengths': ['待 AI 分析后填写'],
                'areas_for_improvement': ['待 AI 分析后填写'],
                'evidence': ['待 AI 分析后填写']
            }

        return analysis

    def _generate_markdown_report(self, analysis, dimensions):
        """生成 Markdown 格式报告"""
        report = f"""# 听评课报告（星火个性化版）

## 基本信息

- **授课教师**: {self.teacher_name}
- **课程名称**: {self.course_name}
- **授课年级**: {self.grade}
- **听课日期**: {self.evaluation_date}
- **录音时长**: {analysis['basic_info']['duration']:.1f} 秒
- **转写字数**: {analysis['basic_info']['word_count']} 字

---

## 一、整体评价

**综合得分**: {analysis['overall_score'] or '待评定'} / 100 分

### 教学亮点

"""
        for highlight in analysis.get('highlights', []):
            report += f"- {highlight}\n"

        report += "\n### 改进建议\n\n"
        for suggestion in analysis.get('suggestions', []):
            report += f"- {suggestion}\n"

        report += "\n---\n\n## 二、分维度评估\n\n"

        for dim_name, dim_info in dimensions.items():
            dim_analysis = analysis['dimensions'].get(dim_name, {})
            items = self._iter_dimension_items(dim_info)
            report += f"""### {dim_name}

**满分**: {dim_info['weight']} 分
**说明**: {dim_info['description']}

"""
            if dim_info.get('sub_items'):
                # 子项明细
                for item in items:
                    sub_score = dim_analysis.get('sub_scores', {}).get(item['name'], '待评定')
                    report += f"""#### {item['name']}（{item['max_score']} 分）

- 得分: {sub_score} 分

"""
                    for indicator in item['indicators']:
                        report += f"- [ ] {indicator}\n"
                    report += "\n"
            else:
                report += "#### 评估指标\n\n"
                for indicator in items[0]['indicators']:
                    report += f"- [ ] {indicator}\n"

            report += f"""
#### 维度得分: {dim_analysis.get('score', '待评定')} / {dim_info['weight']} 分

#### 优点

"""
            for strength in dim_analysis.get('strengths', []):
                report += f"- {strength}\n"

            report += "\n#### 待改进\n\n"
            for area in dim_analysis.get('areas_for_improvement', []):
                report += f"- {area}\n"

            report += "\n#### 证据引用\n\n"
            for evidence in dim_analysis.get('evidence', []):
                report += f"> {evidence}\n"

            report += "\n---\n\n"

        # 添加附录
        report += """## 附录：课堂实录（节选）

> 以下为课堂录音转写文本节选，供详细分析参考：

```
"""
        transcript_text = self.transcript.get('text', '')
        excerpt = transcript_text[:500] + '...' if len(transcript_text) > 500 else transcript_text
        report += excerpt
        report += "\n```\n\n---\n\n*报告生成时间: " + datetime.now().strftime('%Y-%m-%d %H:%M:%S') + "*\n"

        return report

    def _generate_json_report(self, analysis, dimensions):
        """生成 JSON 格式报告"""
        report_data = {
            'meta': {
                'teacher_name': self.teacher_name,
                'course_name': self.course_name,
                'grade': self.grade,
                'evaluation_date': self.evaluation_date,
                'generated_at': datetime.now().isoformat(),
                'standard': '星火教育 1对1 课堂听评课评价标准'
            },
            'transcript_info': analysis['basic_info'],
            'evaluation': {
                'overall_score': analysis['overall_score'],
                'dimensions': analysis['dimensions'],
                'highlights': analysis['highlights'],
                'suggestions': analysis['suggestions']
            },
            'criteria': dimensions
        }
        return json.dumps(report_data, ensure_ascii=False, indent=2)

    def _generate_text_report(self, analysis, dimensions):
        """生成纯文本格式报告"""
        lines = [
            "=" * 60,
            "听评课报告（星火个性化版）".center(46),
            "=" * 60,
            "",
            "【基本信息】",
            f"授课教师: {self.teacher_name}",
            f"课程名称: {self.course_name}",
            f"授课年级: {self.grade}",
            f"听课日期: {self.evaluation_date}",
            f"录音时长: {analysis['basic_info']['duration']:.1f} 秒",
            f"转写字数: {analysis['basic_info']['word_count']} 字",
            "",
            "-" * 60,
            "一、整体评价",
            f"综合得分: {analysis['overall_score'] or '待评定'} / 100 分",
            "",
            "【教学亮点】",
        ]

        for highlight in analysis.get('highlights', []):
            lines.append(f"  + {highlight}")

        lines.extend([
            "",
            "【改进建议】",
        ])

        for suggestion in analysis.get('suggestions', []):
            lines.append(f"  + {suggestion}")

        lines.extend(["", "-" * 60, "二、分维度评估", ""])

        for dim_name, dim_info in dimensions.items():
            dim_analysis = analysis['dimensions'].get(dim_name, {})
            lines.extend([
                f"【{dim_name}】(满分: {dim_info['weight']} 分)",
                f"说明: {dim_info['description']}",
                "",
            ])
            for item in self._iter_dimension_items(dim_info):
                sub_score = dim_analysis.get('sub_scores', {}).get(item['name'], '待评定')
                lines.append(f"  [{item['name']}] ({item['max_score']} 分) - 得分: {sub_score}")
                for indicator in item['indicators']:
                    lines.append(f"    □ {indicator}")

            lines.extend([
                "",
                f"维度得分: {dim_analysis.get('score', '待评定')} / {dim_info['weight']} 分",
                "",
                "优点:",
            ])
            for strength in dim_analysis.get('strengths', []):
                lines.append(f"  + {strength}")

            lines.extend(["", "待改进:"])
            for area in dim_analysis.get('areas_for_improvement', []):
                lines.append(f"  + {area}")

            lines.append("")

        lines.extend([
            "-" * 60,
            f"报告生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
            "=" * 60,
        ])

        return "\n".join(lines)

    def _generate_html_report(self, analysis, dimensions):
        """生成 HTML 格式报告"""
        html = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>听评课报告（星火个性化版） - {self.teacher_name}</title>
    <style>
        body {{
            font-family: "Microsoft YaHei", "SimSun", sans-serif;
            max-width: 900px;
            margin: 40px auto;
            padding: 20px;
            line-height: 1.8;
            color: #333;
        }}
        h1 {{
            color: #2c3e50;
            border-bottom: 3px solid #3498db;
            padding-bottom: 15px;
        }}
        h2 {{
            color: #34495e;
            margin-top: 30px;
        }}
        h3 {{
            color: #555;
        }}
        .info-box {{
            background: #ecf0f1;
            padding: 15px;
            border-radius: 5px;
            margin: 20px 0;
        }}
        .score {{
            font-size: 24px;
            font-weight: bold;
            color: #27ae60;
        }}
        .dimension {{
            background: #fff;
            border-left: 4px solid #3498db;
            padding: 15px;
            margin: 15px 0;
        }}
        .sub-item {{
            background: #f8f9fa;
            border: 1px solid #e0e0e0;
            border-radius: 4px;
            padding: 10px 15px;
            margin: 10px 0;
        }}
        .strength {{
            color: #27ae60;
        }}
        .improvement {{
            color: #e74c3c;
        }}
        blockquote {{
            background: #f9f9f9;
            border-left: 4px solid #ddd;
            padding: 10px 20px;
            margin: 10px 0;
        }}
    </style>
</head>
<body>
    <h1>听评课报告（星火个性化版）</h1>

    <div class="info-box">
        <p><strong>授课教师</strong>: {self.teacher_name}</p>
        <p><strong>课程名称</strong>: {self.course_name}</p>
        <p><strong>授课年级</strong>: {self.grade}</p>
        <p><strong>听课日期</strong>: {self.evaluation_date}</p>
        <p><strong>录音时长</strong>: {analysis['basic_info']['duration']:.1f} 秒</p>
        <p><strong>转写字数</strong>: {analysis['basic_info']['word_count']} 字</p>
    </div>

    <h2>一、整体评价</h2>
    <p class="score">综合得分: {analysis['overall_score'] or '待评定'} / 100 分</p>

    <h3>教学亮点</h3>
    <ul>
"""
        for highlight in analysis.get('highlights', []):
            html += f"        <li>{highlight}</li>\n"

        html += """    </ul>

    <h3>改进建议</h3>
    <ul>
"""
        for suggestion in analysis.get('suggestions', []):
            html += f"        <li>{suggestion}</li>\n"

        html += """    </ul>

    <h2>二、分维度评估</h2>
"""

        for dim_name, dim_info in dimensions.items():
            dim_analysis = analysis['dimensions'].get(dim_name, {})
            html += f"""
    <div class="dimension">
        <h3>{dim_name} <span style="color: #888; font-size: 14px;">(满分: {dim_info['weight']} 分)</span></h3>
        <p><strong>{dim_info['description']}</strong></p>
"""
            for item in self._iter_dimension_items(dim_info):
                sub_score = dim_analysis.get('sub_scores', {}).get(item['name'], '待评定')
                html += f"""
        <div class="sub-item">
            <strong>{item['name']}</strong> <span style="color: #888;">({item['max_score']} 分) - 得分: {sub_score}</span>
            <ul>
"""
                for indicator in item['indicators']:
                    html += f"                <li><input type='checkbox' disabled> {indicator}</li>\n"
                html += """            </ul>
        </div>
"""

            html += f"""
        <p><strong>维度得分</strong>: {dim_analysis.get('score', '待评定')} / {dim_info['weight']} 分</p>

        <h4 class="strength">优点</h4>
        <ul>
"""
            for strength in dim_analysis.get('strengths', []):
                html += f"            <li>{strength}</li>\n"

            html += """        </ul>

        <h4 class="improvement">待改进</h4>
        <ul>
"""
            for area in dim_analysis.get('areas_for_improvement', []):
                html += f"            <li>{area}</li>\n"

            html += """        </ul>
    </div>
"""

        html += f"""
    <hr>
    <p style="text-align: center; color: #888; font-size: 12px;">
        报告生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}
    </p>
</body>
</html>
"""
        return html

    # ---------- 文件名与输出目录 ----------

    @staticmethod
    def _safe_filename(name):
        """去除 Windows 文件名非法字符"""
        for ch in '\\/:*?"<>|':
            name = name.replace(ch, '_')
        return name.strip()

    def _report_date_str(self, date=None):
        """确定报告日期：--date 优先，其次转写文件时间，最后今天"""
        if date:
            try:
                return datetime.strptime(date, '%Y-%m-%d').strftime('%Y%m%d')
            except ValueError:
                pass
        created = self.transcript.get('created_at') or self.transcript.get('timestamp') or ''
        if created:
            try:
                return datetime.fromisoformat(str(created).replace('Z', '+00:00')).strftime('%Y%m%d')
            except Exception:
                pass
        return datetime.now().strftime('%Y%m%d')

    def default_output_name(self, fmt, date=None):
        """
        按统一规则生成报告文件名：
        {YYYYMMDD}_{教师名}老师_{学生名}_听评课报告.{ext}
        例：20260822_吴浩老师_刘子榆_听评课报告.docx
        """
        date_str = self._report_date_str(date)
        teacher = self._safe_filename(self.teacher_name) if self.teacher_name != '待填写' else ''
        student = self._safe_filename(self.student_name) if self.student_name != '待填写' else ''

        teacher_part = f'{teacher}老师' if teacher and not teacher.endswith('老师') else teacher
        parts = [p for p in [date_str, teacher_part, student] if p]
        base = '_'.join(parts) + '_听评课报告'
        ext = {'markdown': 'md', 'json': 'json', 'html': 'html', 'text': 'txt',
               'docx': 'docx', 'pdf': 'pdf'}[fmt]
        return f'{base}.{ext}'

    # ---------- Word 导出 ----------

    def _generate_docx_report(self, analysis, dimensions, output_path):
        """生成 Word (.docx) 格式报告（中文字体：微软雅黑）"""
        doc = Document()

        # 默认样式：微软雅黑 10.5pt（含东亚字体）
        normal = doc.styles['Normal']
        normal.font.name = 'Microsoft YaHei'
        normal.font.size = Pt(10.5)
        normal.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

        def add_title(text, size=16, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6):
            p = doc.add_paragraph()
            p.alignment = align
            p.paragraph_format.space_after = Pt(space_after)
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(size)
            return p

        def add_heading(text, size=13, space_before=10):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(space_before)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(size)
            return p

        def add_para(text, indent=0, size=10.5, bold=False):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(indent)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.bold = bold
            r.font.size = Pt(size)
            return p

        def add_checkbox(text, indent=18):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(indent)
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(f'□ {text}')
            r.font.size = Pt(10)
            return p

        # 标题
        add_title('听评课报告（星火个性化版）')
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run('评估标准：《星火教育 1对1 课堂听评课评价标准》')
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        # 基本信息表
        add_heading('一、基本信息')
        info_rows = [
            ['授课教师', self.teacher_name, '课程名称', self.course_name],
            ['授课年级', self.grade, '学生姓名', self.student_name],
            ['听课日期', self.evaluation_date, '录音时长', f"{analysis['basic_info']['duration']:.1f} 秒"],
            ['转写字数', f"{analysis['basic_info']['word_count']} 字", '', ''],
        ]
        table = doc.add_table(rows=len(info_rows), cols=4)
        table.style = 'Table Grid'
        for i, row in enumerate(info_rows):
            for j, cell_text in enumerate(row):
                cell = table.cell(i, j)
                cell.text = ''
                run = cell.paragraphs[0].add_run(str(cell_text))
                run.font.size = Pt(10)
                if j % 2 == 0:
                    run.bold = True

        # 整体评价
        add_heading('二、整体评价')
        add_para(f"综合得分：{analysis['overall_score'] or '待评定'} / 100 分", bold=True)
        add_para('教学亮点：', bold=True)
        for h in analysis.get('highlights', []):
            add_para(f'• {h}', indent=18)
        add_para('改进建议：', bold=True)
        for s in analysis.get('suggestions', []):
            add_para(f'• {s}', indent=18)

        # 分维度评估
        add_heading('三、分维度评估')
        for dim_name, dim_info in dimensions.items():
            dim_analysis = analysis['dimensions'].get(dim_name, {})
            add_heading(f'{dim_name}（满分 {dim_info["weight"]} 分）', size=12)
            add_para(f"说明：{dim_info['description']}", size=9.5)

            for item in self._iter_dimension_items(dim_info):
                sub_score = dim_analysis.get('sub_scores', {}).get(item['name'], '待评定')
                add_para(f"▎{item['name']}（{item['max_score']} 分）— 得分：{sub_score}", bold=True, indent=6)
                for indicator in item['indicators']:
                    add_checkbox(indicator)

            add_para(f"维度得分：{dim_analysis.get('score', '待评定')} / {dim_info['weight']} 分", bold=True, indent=6)
            add_para('优点：', bold=True, indent=6)
            for strength in dim_analysis.get('strengths', []):
                add_para(f'• {strength}', indent=24)
            add_para('待改进：', bold=True, indent=6)
            for area in dim_analysis.get('areas_for_improvement', []):
                add_para(f'• {area}', indent=24)
            add_para('证据引用：', bold=True, indent=6)
            for evidence in dim_analysis.get('evidence', []):
                add_para(f'“{evidence}”', indent=24, size=9.5)

        # 附录
        add_heading('附录：课堂实录（节选）')
        excerpt = self.transcript.get('text', '')[:500]
        add_para(excerpt + ('...' if len(self.transcript.get('text', '')) > 500 else ''), size=9)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f'报告生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}')
        r.font.size = Pt(8)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        print(f"报告已保存: {output_path}")
        return str(output_path)

    # ---------- PDF 导出 ----------

    @staticmethod
    def _register_cjk_font():
        """注册系统中文字体（reportlab），返回字体名"""
        candidates = [
            ('MSYH', r'C:\Windows\Fonts\msyh.ttc', 0),   # 微软雅黑
            ('SimHei', r'C:\Windows\Fonts\simhei.ttf', None),  # 黑体
            ('SimSun', r'C:\Windows\Fonts\simsun.ttc', 0),  # 宋体
            ('MSYH', '/usr/share/fonts/truetype/wqy/wqy-microhei.ttc', None),  # Linux WQY
            ('SimHei', '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc', None),  # Linux WQY
            ('SimSun', '/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc', 0),  # Linux Noto
        ]
        for name, path, subfont in candidates:
            try:
                if subfont is not None:
                    pdfmetrics.registerFont(TTFont(name, path, subfontIndex=subfont))
                else:
                    pdfmetrics.registerFont(TTFont(name, path))
                return name
            except Exception:
                continue
        raise RuntimeError('未找到可用的中文字体（msyh.ttc / simhei.ttf / simsun.ttc）')

    def _generate_pdf_report(self, analysis, dimensions, output_path):
        """生成 PDF 格式报告（reportlab + 系统中文字体）"""
        font_name = self._register_cjk_font()

        title_style = ParagraphStyle('TitleCN', fontName=font_name, fontSize=17, leading=24,
                                     alignment=1, spaceAfter=4)
        sub_style = ParagraphStyle('SubCN', fontName=font_name, fontSize=9, leading=13,
                                   alignment=1, textColor=colors.HexColor('#888888'), spaceAfter=10)
        h1_style = ParagraphStyle('H1CN', fontName=font_name, fontSize=13, leading=19,
                                  spaceBefore=10, spaceAfter=4)
        body_style = ParagraphStyle('BodyCN', fontName=font_name, fontSize=10, leading=15,
                                    wordWrap='CJK')
        small_style = ParagraphStyle('SmallCN', fontName=font_name, fontSize=9, leading=13,
                                     wordWrap='CJK')
        bold_style = ParagraphStyle('BoldCN', fontName=font_name, fontSize=10.5, leading=15,
                                    wordWrap='CJK')
        center_small = ParagraphStyle('CenterSmall', fontName=font_name, fontSize=8, leading=11,
                                      alignment=1, textColor=colors.HexColor('#888888'))

        def esc(text):
            return str(text).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

        story = []
        story.append(Paragraph('听评课报告（星火个性化版）', title_style))
        story.append(Paragraph('评估标准：《星火教育 1对1 课堂听评课评价标准》', sub_style))

        # 基本信息表
        story.append(Paragraph('一、基本信息', h1_style))
        info_rows = [
            ['授课教师', self.teacher_name, '课程名称', self.course_name],
            ['授课年级', self.grade, '学生姓名', self.student_name],
            ['听课日期', self.evaluation_date, '录音时长', f"{analysis['basic_info']['duration']:.1f} 秒"],
            ['转写字数', f"{analysis['basic_info']['word_count']} 字", '', ''],
        ]
        data = [[Paragraph(f'<b>{esc(c)}</b>' if j % 2 == 0 else esc(c), small_style) for j, c in enumerate(row)]
                for row in info_rows]
        t = Table(data, colWidths=[3 * cm, 4.5 * cm, 3 * cm, 4.5 * cm])
        t.setStyle(TableStyle([
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#BBBBBB')),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ]))
        story.append(t)

        # 整体评价
        story.append(Paragraph('二、整体评价', h1_style))
        story.append(Paragraph(f"综合得分：<b>{analysis['overall_score'] or '待评定'} / 100 分</b>", bold_style))
        story.append(Paragraph('教学亮点：', bold_style))
        for h in analysis.get('highlights', []):
            story.append(Paragraph(f'• {esc(h)}', body_style))
        story.append(Paragraph('改进建议：', bold_style))
        for s in analysis.get('suggestions', []):
            story.append(Paragraph(f'• {esc(s)}', body_style))

        # 分维度评估
        story.append(Paragraph('三、分维度评估', h1_style))
        for dim_name, dim_info in dimensions.items():
            dim_analysis = analysis['dimensions'].get(dim_name, {})
            story.append(Paragraph(f'{esc(dim_name)}（满分 {dim_info["weight"]} 分）', h1_style))
            story.append(Paragraph(f"说明：{esc(dim_info['description'])}", small_style))

            for item in self._iter_dimension_items(dim_info):
                sub_score = dim_analysis.get('sub_scores', {}).get(item['name'], '待评定')
                story.append(Paragraph(
                    f"▎{esc(item['name'])}（{item['max_score']} 分）— 得分：{sub_score}", bold_style))
                for indicator in item['indicators']:
                    story.append(Paragraph(f'□ {esc(indicator)}', body_style))

            story.append(Paragraph(
                f"维度得分：<b>{dim_analysis.get('score', '待评定')} / {dim_info['weight']} 分</b>", bold_style))
            story.append(Paragraph('优点：', bold_style))
            for strength in dim_analysis.get('strengths', []):
                story.append(Paragraph(f'• {esc(strength)}', body_style))
            story.append(Paragraph('待改进：', bold_style))
            for area in dim_analysis.get('areas_for_improvement', []):
                story.append(Paragraph(f'• {esc(area)}', body_style))
            story.append(Paragraph('证据引用：', bold_style))
            for evidence in dim_analysis.get('evidence', []):
                story.append(Paragraph(f'“{esc(evidence)}”', small_style))

        # 附录
        story.append(Paragraph('附录：课堂实录（节选）', h1_style))
        excerpt = self.transcript.get('text', '')[:500]
        story.append(Paragraph(esc(excerpt) + ('...' if len(self.transcript.get('text', '')) > 500 else ''), small_style))

        story.append(Spacer(1, 16))
        story.append(Paragraph(f'报告生成时间：{datetime.now().strftime("%Y-%m-%d %H:%M:%S")}', center_small))

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        SimpleDocTemplate(str(output_path), pagesize=A4,
                          leftMargin=2 * cm, rightMargin=2 * cm,
                          topMargin=2 * cm, bottomMargin=2 * cm,
                          title='听评课报告（星火个性化版）').build(story)
        print(f"报告已保存: {output_path}")
        return str(output_path)


def main():
    """命令行入口"""
    import argparse

    parser = argparse.ArgumentParser(description='听评课报告生成工具（星火个性化版）')
    parser.add_argument('transcript', help='转写文件路径 (JSON 格式)')
    parser.add_argument('-o', '--output', help='输出文件路径')
    parser.add_argument('-f', '--format', default='markdown',
                       choices=['markdown', 'json', 'html', 'text', 'docx', 'pdf'],
                       help='输出格式 (默认: markdown，Word/PDF 输出到统一报告目录)', )
    parser.add_argument('--teacher', help='授课教师姓名（用于报告命名，如：吴浩）')
    parser.add_argument('--course', help='课程名称')
    parser.add_argument('--grade', help='授课年级')
    parser.add_argument('--student', help='学生姓名（用于报告命名，如：刘子榆）')
    parser.add_argument('--date', help='听课日期 (YYYY-MM-DD)，默认取转写日期或今天')
    parser.add_argument('--output-dir', default=None,
                        help=f'报告输出目录（默认统一目录: {DEFAULT_REPORT_DIR}）')
    parser.add_argument('--standard', help='自定义评估标准 JSON 文件路径（可选）')

    args = parser.parse_args()

    # 自定义标准支持
    custom_template = None
    if args.standard:
        with open(args.standard, 'r', encoding='utf-8') as f:
            custom_template = json.load(f)

    # 生成报告
    reporter = ClassroomEvaluationReport(
        args.transcript,
        teacher_name=args.teacher,
        course_name=args.course,
        grade=args.grade,
        student_name=args.student
    )

    output_path = args.output
    if output_path is None:
        # 统一输出目录 + 统一命名：{YYYYMMDD}_{教师}老师_{学生}_听评课报告.{ext}
        out_dir = args.output_dir or DEFAULT_REPORT_DIR
        output_path = str(Path(out_dir) / reporter.default_output_name(args.format, date=args.date))

    report = reporter.generate_report(format=args.format, output_path=output_path, custom_template=custom_template)
    print(f"\n报告生成完成！")


if __name__ == '__main__':
    main()
