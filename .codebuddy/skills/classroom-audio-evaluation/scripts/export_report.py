# -*- coding: utf-8 -*-
"""
听评课报告导出脚本：Markdown → PDF / Word(.docx)

特性：
- 支持数学/物理等学科公式（LaTeX 语法，行内 $...$ 与块级 $$...$$）
- 表格单元格内的 **加粗** 与行内公式均正确渲染（不残留星号）
- PDF 排版美化：主题配色、分区标题、页眉页脚、表头底色与斑马纹

用法示例：
    python export_report.py 报告.md --teacher 何东奇 --student 何名慧 --date 2026-08-22 -f pdf
    python export_report.py 报告.md -f docx

输出：默认目录 ./reports/，命名 {YYYYMMDD}_{教师}老师_{学生}_听评课报告.{pdf|docx}
"""

import os
import re
import sys
import tempfile
from pathlib import Path
from datetime import datetime

# Windows GBK 控制台兼容
for _stream in (sys.stdout, sys.stderr):
    try:
        _stream.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass

DEFAULT_REPORT_DIR = './reports'

# 主题配色
COLOR_PRIMARY = '#1F4E79'      # 深蓝（标题）
COLOR_SECONDARY = '#2E74B5'    # 中蓝（次级标题）
COLOR_TEXT = '#222222'
COLOR_MUTED = '#666666'
COLOR_HEADER_BG = '#EEF3FA'    # 表头底色
COLOR_ROW_ALT = '#F7FAFD'      # 斑马纹


# ---------------- 数学公式渲染（matplotlib mathtext） ----------------
import matplotlib
matplotlib.use('Agg')

# ---------------- 修复 reportlab 5.0.x CJK 断行 bug ----------------
# 当 wordWrap='CJK' 的段落内联 <img> 公式时，reportlab 的 cjkFragSplit 会对
# 空文本的 img fragment 调用 ord('')，抛出 TypeError。这里 monkey-patch 修复之。
import reportlab.platypus.paragraph as _rlp
from unicodedata import category as _category
from reportlab.rl_config import _FUZZ as _FUZZ
from reportlab.lib.utils import isBytes as _isBytes


def _cjk_frag_split_fixed(frags, maxWidths, calcBounds, encoding='utf8'):
    U = []
    for f in frags:
        text = f.text
        if _isBytes(text):
            text = text.decode(encoding)
        if text:
            U.extend([_rlp.cjkU(t, f, encoding) for t in text])
        else:
            U.append(_rlp.cjkU(text, f, encoding))
    lines = []
    i = widthUsed = lineStartPos = 0
    maxWidth = maxWidths[0]
    nU = len(U)
    while i < nU:
        u = U[i]
        i += 1
        w = u.width
        if hasattr(w, 'normalizedValue'):
            w._normalizer = maxWidth
            w = w.normalizedValue(maxWidth)
        widthUsed += w
        lineBreak = hasattr(u.frag, 'lineBreak')
        endLine = (widthUsed > maxWidth + _FUZZ and widthUsed > 0) or lineBreak
        if endLine:
            extraSpace = maxWidth - widthUsed
            if not lineBreak:
                if u and ord(u) < 0x3000:
                    limitCheck = (lineStartPos + i) >> 1
                    for j in range(i - 1, limitCheck, -1):
                        uj = U[j]
                        if uj and (_category(uj) == 'Zs' or ord(uj) >= 0x3000):
                            k = j + 1
                            if k < i:
                                j = k + 1
                                extraSpace += sum(U[ii].width for ii in range(j, i))
                                w = U[k].width
                                u = U[k]
                                i = j
                                break
                if u not in _rlp.ALL_CANNOT_START and i > lineStartPos + 1:
                    i -= 1
                    extraSpace += w
            lines.append(_rlp.makeCJKParaLine(U[lineStartPos:i], maxWidth, widthUsed, extraSpace, lineBreak, calcBounds))
            try:
                maxWidth = maxWidths[len(lines)]
            except IndexError:
                maxWidth = maxWidths[-1]
            lineStartPos = i
            widthUsed = 0
    if widthUsed > 0:
        lines.append(_rlp.makeCJKParaLine(U[lineStartPos:], maxWidth, widthUsed, maxWidth - widthUsed, False, calcBounds))
    return _rlp.ParaLines(kind=1, lines=lines)


_rlp.cjkFragSplit = _cjk_frag_split_fixed


# 补全中文全角标点禁则（不能出现在行首的标点）。
# reportlab 默认 ALL_CANNOT_START 缺少全角逗号/分号/冒号/问号/感叹号/引号等，
# 导致这些标点被错误地放到行首（排版"标点行首"、上一行未满即换行）。
_rlp.ALL_CANNOT_START = _rlp.ALL_CANNOT_START + (
    '\uff0c'   # ，
    '\uff1b'   # ；
    '\uff1a'   # ：
    '\uff1f'   # ？
    '\uff01'   # ！
    '\uff0e'   # ．
    '\u201d'   # ”
    '\u2019'   # ’
    '\u2026'   # …
    '\u3009'   # 〉
    '\u300b'   # 》
)


def _break_lines_cjk_fixed(self, maxWidths):
    """让 wordWrap='CJK' 的段落（含单 frag 纯文本）统一走 cjkFragSplit。

    reportlab 原始 breakLinesCJK 对「单一 frag、无图片」的段落走 wordSplit，
    而 wordSplit 不处理中文标点禁则，导致评语里出现「标点行首、未满即换行」。
    """
    if not isinstance(maxWidths, (list, tuple)):
        maxWidths = [maxWidths]
    style = self.style
    self.height = 0
    _rlp._handleBulletWidth(self.bulletText, style, maxWidths)
    frags = self.frags
    nFrags = len(frags)
    if nFrags <= 0:
        return _rlp.ParaLines(kind=0, fontSize=style.fontSize, fontName=style.fontName,
                              textColor=style.textColor, lines=[],
                              ascent=style.fontSize, descent=-0.2 * style.fontSize)
    if hasattr(self, 'blPara') and getattr(self, '_splitpara', 0):
        return self.blPara
    autoLeading = getattr(self, 'autoLeading', getattr(style, 'autoLeading', ''))
    calcBounds = autoLeading not in ('', 'off')
    return _rlp.cjkFragSplit(frags, maxWidths, calcBounds)


_rlp.Paragraph.breakLinesCJK = _break_lines_cjk_fixed

_formula_dir = None
_formula_counter = [0]


def _get_formula_dir():
    global _formula_dir
    if _formula_dir is None:
        _formula_dir = tempfile.mkdtemp(prefix='tmt_formula_')
    return _formula_dir


def render_math(tex, fontsize=11, dpi=220, color='#1a1a1a'):
    """渲染 LaTeX 公式为 PNG，返回 (路径, 宽度pt, 高度pt)。失败抛出异常。"""
    import matplotlib.pyplot as plt
    from PIL import Image

    _get_formula_dir()
    _formula_counter[0] += 1
    path = os.path.join(_formula_dir, f'f{_formula_counter[0]}.png')

    fig = plt.figure()
    fig.text(0, 0, f'${tex}$', fontsize=fontsize, color=color)
    try:
        fig.savefig(path, dpi=dpi, transparent=True,
                    bbox_inches='tight', pad_inches=0.03)
    finally:
        plt.close(fig)

    with Image.open(path) as im:
        w_px, h_px = im.size
    scale = 72.0 / dpi
    return path, w_px * scale, h_px * scale


# ---------------- Markdown 解析 ----------------

def parse_markdown(md_text):
    """解析为块列表：('h1'|'h2'|'h3'|'h4'|'para'|'bullet'|'quote'|'table'|'hr'|'formula', 内容)"""
    blocks = []
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue
        if stripped == '---':
            blocks.append(('hr', ''))
            i += 1
            continue
        # 块级公式 $$...$$
        if stripped.startswith('$$'):
            tex = stripped[2:].strip()
            if tex.endswith('$$'):
                tex = tex[:-2].strip()
                blocks.append(('formula', tex))
                i += 1
                continue
            buf = [tex]
            i += 1
            while i < len(lines):
                ln = lines[i].strip()
                if ln.endswith('$$'):
                    buf.append(ln[:-2].strip())
                    i += 1
                    break
                buf.append(ln)
                i += 1
            blocks.append(('formula', ' '.join(buf)))
            continue
        m = re.match(r'^(#{1,4})\s+(.*)$', stripped)
        if m:
            level = len(m.group(1))
            blocks.append((f'h{level}', m.group(2).strip()))
            i += 1
            continue
        if stripped.startswith('```'):
            code_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code_lines.append(lines[i])
                i += 1
            i += 1
            blocks.append(('code', '\n'.join(code_lines)))
            continue
        if stripped.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{2,}:?', c) for c in cells if c.strip()):
                    rows.append(cells)
                i += 1
            blocks.append(('table', rows))
            continue
        if stripped.startswith('> '):
            blocks.append(('quote', stripped[2:].strip()))
            i += 1
            continue
        m = re.match(r'^[-*]\s+(.*)$', stripped)
        if m:
            blocks.append(('bullet', m.group(1).strip()))
            i += 1
            continue
        m = re.match(r'^\d+[.、]\s+(.*)$', stripped)
        if m:
            blocks.append(('bullet', m.group(1).strip()))
            i += 1
            continue
        blocks.append(('para', stripped))
        i += 1
    return blocks


def split_bold(text):
    """把 **加粗** 拆成 [(text, bold), ...]"""
    parts = re.split(r'(\*\*.*?\*\*)', text)
    result = []
    for p in parts:
        if not p:
            continue
        if p.startswith('**') and p.endswith('**') and len(p) > 4:
            result.append((p[2:-2], True))
        else:
            result.append((p, False))
    return result


# ---------------- 通用内联处理 ----------------

def esc(text):
    return text.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')


def _strip_inline_math(text):
    """移除行内 $...$，仅保留纯文本（用于 docx 等不支持公式的场景）。"""
    return re.sub(r'\$[^$]*\$', lambda m: m.group(0)[1:-1], text)


def rich_inline_pdf(text, fontsize=10):
    """将行内公式 $...$ 与 **加粗** 转为 reportlab Paragraph 标记。"""
    parts = re.split(r'(\$[^$]*\$)', text)
    out = []
    for p in parts:
        if not p:
            continue
        if p.startswith('$') and p.endswith('$') and len(p) > 2:
            tex = p[1:-1].strip()
            try:
                path, w, h = render_math(tex, fontsize=fontsize + 1)
                out.append(f'<img src="{path}" width="{w:.1f}" height="{h:.1f}" valign="middle"/>')
            except Exception:
                out.append(esc(tex))
        else:
            for s, bold in split_bold(p):
                out.append(f'<b>{esc(s)}</b>' if bold else esc(s))
    return _wrap_symbols(''.join(out))


# ---------------- Word 导出 ----------------

def md_to_docx(md_text, output_path):
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Microsoft YaHei'
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

    def add_runs(p, text, size=None, base_bold=False, color=None):
        for seg, bold in split_bold(_strip_inline_math(text)):
            r = p.add_run(seg)
            if size:
                r.font.size = Pt(size)
            r.bold = base_bold or bold
            if color:
                r.font.color.rgb = color

    def add_heading(text, level):
        p = doc.add_paragraph()
        if level == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            add_runs(p, text, size=17, base_bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
        elif level == 2:
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            add_runs(p, text, size=14, base_bold=True, color=RGBColor(0x1F, 0x4E, 0x79))
        elif level == 3:
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, text, size=12, base_bold=True, color=RGBColor(0x2E, 0x74, 0xB5))
        else:
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, text, size=11, base_bold=True)

    blocks = parse_markdown(md_text)
    for kind, content in blocks:
        if kind == 'hr':
            continue
        elif kind in ('h1', 'h2', 'h3', 'h4'):
            add_heading(content, int(kind[1]))
        elif kind == 'formula':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, _strip_inline_math(content), size=11)
        elif kind == 'bullet':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, '• ' + _strip_inline_math(content), size=10.5)
        elif kind == 'quote':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, '“' + _strip_inline_math(content) + '”', size=9.5,
                     color=RGBColor(0x66, 0x66, 0x66))
        elif kind == 'para':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, _strip_inline_math(content), size=10.5)
        elif kind == 'code':
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(12)
            add_runs(p, content, size=9, color=RGBColor(0x44, 0x44, 0x44))
        elif kind == 'table':
            if not content:
                continue
            rows = content
            table = doc.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = 'Table Grid'
            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row):
                    if ci >= len(row):
                        continue
                    cell = table.cell(ri, ci)
                    cell.text = ''
                    for seg, bold in split_bold(_strip_inline_math(cell_text)):
                        r = cell.paragraphs[0].add_run(seg)
                        r.bold = bold or (ri == 0)
                        r.font.size = Pt(9.5)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return str(output_path)


# ---------------- PDF 导出 ----------------

def _register_cjk_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [
        ('MSYH', r'C:\Windows\Fonts\msyh.ttc', 0),
        ('SimHei', r'C:\Windows\Fonts\simhei.ttf', None),
        ('SimSun', r'C:\Windows\Fonts\simsun.ttc', 0),
        ('MSYH', '/usr/share/fonts/truetype/wqy/wqy-microhei.ttc', None),
        ('SimHei', '/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc', None),
        ('SimSun', '/usr/share/fonts/opentype/noto/NotoSerifCJK-Regular.ttc', 0),
    ]
    for name, path, subfont in candidates:
        try:
            if subfont is not None:
                pdfmetrics.registerFont(TTFont(name, path, subfontIndex=subfont))
            else:
                pdfmetrics.registerFont(TTFont(name, path))
            return name
        except Exception:
            continue
    raise RuntimeError('未找到可用的中文字体')


_SYMBOL_FONT = None

# 中文字体（微软雅黑/黑体/宋体）缺字形、需用符号字体渲染的字符：
# ✅ ❌ ✓ ✔ ✗ ✘ ☑ ☐ ⭕ ❎ 等 emoji / 对错号
_SYMBOL_RE = re.compile(
    r'[\u2705\u274c\u274e\u2713\u2714\u2717\u2718\u2611\u2610\u2b55]+'
)


def _register_symbol_font():
    """注册 Segoe UI Symbol 符号字体，用于渲染中文字体缺失的 emoji 勾叉。

    Windows 下 ✅(U+2705)/❌(U+274C)/✓(U+2713)/✗(U+2717) 在微软雅黑/黑体/宋体
    中均无字形，直接渲染会变成空白；需改用 Segoe UI Symbol。
    """
    global _SYMBOL_FONT
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    candidates = [
        r'C:\Windows\Fonts\seguisym.ttf',
        '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf',
    ]
    for path in candidates:
        try:
            pdfmetrics.registerFont(TTFont('SegoeUISymbol', path))
            _SYMBOL_FONT = 'SegoeUISymbol'
            return _SYMBOL_FONT
        except Exception:
            continue
    return None


def _wrap_symbols(text):
    """把中文字体缺字形的符号字符包上符号字体，避免 PDF 渲染成空白。"""
    if _SYMBOL_FONT is None:
        return text
    return _SYMBOL_RE.sub(
        lambda m: f'<font name="{_SYMBOL_FONT}">{m.group(0)}</font>', text
    )


def _build_pdf_styles(font_name):
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib import colors
    return {
        'h1': ParagraphStyle('h1', fontName=font_name, fontSize=18, leading=26,
                             alignment=1, spaceAfter=12, textColor=colors.HexColor(COLOR_PRIMARY)),
        'h2': ParagraphStyle('h2', fontName=font_name, fontSize=13, leading=20,
                             spaceBefore=14, spaceAfter=4, textColor=colors.HexColor(COLOR_PRIMARY)),
        'h3': ParagraphStyle('h3', fontName=font_name, fontSize=11.5, leading=17,
                             spaceBefore=8, spaceAfter=3, textColor=colors.HexColor(COLOR_SECONDARY)),
        'h4': ParagraphStyle('h4', fontName=font_name, fontSize=10.5, leading=16,
                             spaceBefore=6, spaceAfter=2, textColor=colors.HexColor(COLOR_SECONDARY)),
        'body': ParagraphStyle('body', fontName=font_name, fontSize=10, leading=16,
                               wordWrap='CJK', spaceAfter=4, textColor=colors.HexColor(COLOR_TEXT)),
        'bullet': ParagraphStyle('bullet', fontName=font_name, fontSize=10, leading=16,
                                 wordWrap='CJK', leftIndent=14, spaceAfter=3,
                                 textColor=colors.HexColor(COLOR_TEXT)),
        'quote': ParagraphStyle('quote', fontName=font_name, fontSize=9.5, leading=14,
                                wordWrap='CJK', leftIndent=14,
                                textColor=colors.HexColor(COLOR_MUTED), spaceAfter=2),
        'code': ParagraphStyle('code', fontName=font_name, fontSize=8.5, leading=12,
                               wordWrap='CJK', leftIndent=10,
                               textColor=colors.HexColor('#444444')),
        'cell_left': ParagraphStyle('cell_left', fontName=font_name, fontSize=9.5, leading=14,
                                    wordWrap='CJK', alignment=0,
                                    textColor=colors.HexColor(COLOR_TEXT)),
        'cell_center': ParagraphStyle('cell_center', fontName=font_name, fontSize=9.5, leading=14,
                                      wordWrap='CJK', alignment=1,
                                      textColor=colors.HexColor(COLOR_TEXT)),
    }


def _table_style(col_widths):
    from reportlab.lib import colors
    from reportlab.platypus import TableStyle
    return TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor(COLOR_HEADER_BG)),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1),
         [colors.white, colors.HexColor(COLOR_ROW_ALT)]),
        ('GRID', (0, 0), (-1, -1), 0.4, colors.HexColor('#D5DEE8')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ])


def md_to_pdf(md_text, output_path):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                    Image, HRFlowable)

    font_name = _register_cjk_font()
    _register_symbol_font()
    styles = _build_pdf_styles(font_name)

    def h2_rule():
        return HRFlowable(width='100%', thickness=0.8,
                          color=colors.HexColor(COLOR_SECONDARY),
                          spaceBefore=2, spaceAfter=6)

    story = []
    blocks = parse_markdown(md_text)
    for kind, content in blocks:
        if kind == 'hr':
            story.append(Spacer(1, 6))
        elif kind in ('h1', 'h2', 'h3', 'h4'):
            story.append(Paragraph(rich_inline_pdf(content, fontsize=12), styles[kind]))
            if kind == 'h2':
                story.append(h2_rule())
        elif kind == 'formula':
            try:
                path, w, h = render_math(content, fontsize=12)
                story.append(Spacer(1, 4))
                story.append(Image(path, width=w, height=h))
                story.append(Spacer(1, 4))
            except Exception:
                story.append(Paragraph(esc(content), styles['body']))
        elif kind == 'bullet':
            story.append(Paragraph(f'• {rich_inline_pdf(content)}', styles['bullet']))
        elif kind == 'quote':
            story.append(Paragraph(f'“{rich_inline_pdf(content)}”', styles['quote']))
        elif kind == 'para':
            story.append(Paragraph(rich_inline_pdf(content), styles['body']))
        elif kind == 'code':
            story.append(Paragraph(esc(content), styles['code']))
        elif kind == 'table' and content:
            rows = content
            ncols = max(len(r) for r in rows)
            total_w = 17 * cm

            def _vis_w(cell):
                # 估算单元格内容视觉宽度：中文算 2、ASCII 算 1，公式按 8 字符宽估算
                plain = re.sub(r'\$[^$]*\$', 'XXXXXXXX', cell)
                plain = re.sub(r'\*\*([^*]+)\*\*', r'\1', plain)
                plain = re.sub(r'`([^`]+)`', r'\1', plain)
                plain = re.sub(r'<[^>]+>', '', plain)
                return sum(2 if ord(ch) > 127 else 1 for ch in plain)

            # 每列对齐：短内容（视觉宽 <=14，约 7 个汉字）居中，长描述（简评/评述/评价）靠左
            aligns = []
            for c in range(ncols):
                max_w = max((_vis_w(r[c]) for r in rows[1:] if c < len(r)), default=0)
                aligns.append('center' if max_w <= 14 else 'left')

            # 列宽分配：短列（数值/标签）保证最宽内容单行、长列（描述）允许换行，
            # 剩余空间按内容视觉权重分配给长列；避免短数值列被挤成竖排
            UNIT = 4.75   # 一个视觉宽单位 ≈ 0.5 个 9.5pt 字符的宽度
            PAD = 14      # 左右 padding 合计(12pt) + 缓冲(2pt)
            SHORT = 20    # 视觉宽 <= 此值视为短列，须保证单行
            col_max = []
            for c in range(ncols):
                col_max.append(max((_vis_w(r[c]) for r in rows if c < len(r)),
                                   default=4))
            min_pt = []
            is_short = []
            for c in range(ncols):
                mw = col_max[c]
                if mw <= SHORT:
                    min_pt.append(max(mw * UNIT + PAD, 20))
                    is_short.append(True)
                else:
                    min_pt.append(40)  # 长列允许换行，给较小保底
                    is_short.append(False)
            total_min = sum(min_pt)
            if total_min <= total_w:
                extra = total_w - total_min
                total_wt = sum(max(mw, 4) for mw in col_max)
                col_widths = [min_pt[c] + extra * max(col_max[c], 4) / total_wt
                              for c in range(ncols)]
            else:
                # 空间不足：短列最小宽度为硬约束，剩余空间按权重分给长列
                fixed = sum(min_pt[c] for c in range(ncols) if is_short[c])
                long_idx = [c for c in range(ncols) if not is_short[c]]
                col_widths = list(min_pt)
                if long_idx and fixed < total_w:
                    remain = total_w - fixed
                    lw = sum(max(col_max[c], 4) for c in long_idx)
                    for c in long_idx:
                        col_widths[c] = max(remain * max(col_max[c], 4) / lw, 30)
                else:
                    # 短列本身都放不下，等比压缩
                    col_widths = [w * total_w / total_min for w in min_pt]

            data = []
            for ri, row in enumerate(rows):
                padded = row + [''] * (ncols - len(row))
                if ri == 0:
                    data.append([Paragraph(f'<b>{rich_inline_pdf(c, fontsize=9.5)}</b>',
                                           styles['cell_center']) for c in padded])
                else:
                    data.append([Paragraph(rich_inline_pdf(c, fontsize=9.5),
                                           styles['cell_center' if aligns[ci] == 'center' else 'cell_left'])
                                 for ci, c in enumerate(padded)])
            t = Table(data, colWidths=col_widths, repeatRows=1)
            t.setStyle(_table_style(col_widths))
            story.append(t)
            story.append(Spacer(1, 6))

    def on_page(canvas, doc):
        canvas.saveState()
        # 页眉
        canvas.setFont(font_name, 8)
        canvas.setFillColor(colors.HexColor('#9AA7B5'))
        canvas.drawString(2 * cm, A4[1] - 1.2 * cm, '课堂听评课评估报告')
        canvas.drawRightString(A4[0] - 2 * cm, A4[1] - 1.2 * cm, '星火标准 · 三维度评估')
        # 页脚页码
        canvas.setFont(font_name, 8)
        canvas.drawCentredString(A4[0] / 2, 1.0 * cm, f'第 {doc.page} 页')
        canvas.restoreState()

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = SimpleDocTemplate(str(output_path), pagesize=A4,
                            leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2.2 * cm, bottomMargin=2 * cm,
                            title=Path(output_path).stem)
    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return str(output_path)


# ---------------- 主入口 ----------------

def main():
    import argparse

    parser = argparse.ArgumentParser(description='听评课报告导出：Markdown → PDF/Word')
    parser.add_argument('input', help='Markdown 报告文件路径')
    parser.add_argument('-f', '--format', default='pdf', choices=['docx', 'pdf'],
                        help='输出格式 (默认: pdf)')
    parser.add_argument('--teacher', help='教师姓名（用于命名，如：何东奇）')
    parser.add_argument('--student', help='学生姓名（用于命名，如：何名慧）')
    parser.add_argument('--date', help='听课日期 (YYYY-MM-DD)，默认今天')
    parser.add_argument('--output-dir', default=None,
                        help=f'输出目录（默认: {DEFAULT_REPORT_DIR}）')
    parser.add_argument('-o', '--output', help='指定完整输出文件路径（覆盖命名规则）')
    args = parser.parse_args()

    with open(args.input, 'r', encoding='utf-8') as f:
        md_text = f.read()

    if args.teacher:
        md_text = md_text.replace('（待补充）', args.teacher).replace('待填写', args.teacher)
    if args.student and '学生姓名' not in md_text:
        md_text = md_text.replace('## 基本信息', f'## 基本信息\n\n- **学生姓名**：{args.student}', 1)

    if args.output:
        output_path = args.output
    else:
        if args.date:
            try:
                date_str = datetime.strptime(args.date, '%Y-%m-%d').strftime('%Y%m%d')
            except ValueError:
                date_str = datetime.now().strftime('%Y%m%d')
        else:
            date_str = datetime.now().strftime('%Y%m%d')

        def clean(s):
            for ch in '\\/:*?"<>|':
                s = s.replace(ch, '_')
            return s.strip()

        teacher = clean(args.teacher) if args.teacher else ''
        student = clean(args.student) if args.student else ''
        teacher_part = f'{teacher}老师' if teacher and not teacher.endswith('老师') else teacher
        parts = [p for p in [date_str, teacher_part, student] if p]
        base = '_'.join(parts) + '_听评课报告'
        out_dir = args.output_dir or DEFAULT_REPORT_DIR
        output_path = str(Path(out_dir) / f'{base}.{args.format}')

    if args.format == 'docx':
        saved = md_to_docx(md_text, output_path)
    else:
        saved = md_to_pdf(md_text, output_path)
    print(f'报告已导出: {saved}')


if __name__ == '__main__':
    main()
